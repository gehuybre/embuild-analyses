"""Convert a BBC-DR project workbook without replacing the published source.

Run from analyses/:
    python apps/bouwprojecten-gemeenten/versions/previous/scripts/convert_project_update.py
"""

import argparse
import csv
from decimal import Decimal, InvalidOperation
from pathlib import Path

from openpyxl import load_workbook

from process_project_details import (
    CSV_OUTPUT_COLUMNS,
    DATA_DIR,
    POLICY_COLUMNS,
    extract_code_from_text,
)


TEXT_COLUMNS = [
    "Beleidsdoelst. totaaloverzicht",
    "Actieplan totaaloverzicht",
    "Actie totaaloverzicht",
]
REQUIRED_COLUMNS = ["Bestuur", *TEXT_COLUMNS, "Type rapport", "Rapportjaar", "Boekjaar", "Uitgave"]
REVIEW_COLUMNS = ["NIS-code", "Bestuur", "Actie code", "Reden"]


def text(value: object) -> str:
    return "" if value is None else str(value).strip()


def numeric(value: object, label: str) -> Decimal:
    try:
        number = Decimal(text(value).replace(",", "."))
    except InvalidOperation as error:
        raise ValueError(f"{label}: invalid numeric value {value!r}") from error
    if not number.is_finite():
        raise ValueError(f"{label}: expected a finite number, got {value!r}")
    return number


def integer(value: object, label: str) -> str:
    number = numeric(value, label)
    if number != number.to_integral_value() or number <= 0:
        raise ValueError(f"{label}: expected a positive integer, got {value!r}")
    return str(int(number))


def load_reference(reference: Path) -> tuple[
    dict[str, str], dict[tuple[str, str], set[tuple[str, ...]]]
]:
    names: dict[str, str] = {}
    policies: dict[tuple[str, str], set[tuple[str, ...]]] = {}
    with reference.open(encoding="utf-8-sig", newline="") as handle:
        reader = csv.DictReader(handle, delimiter=";")
        required = {"Bestuur", "NIS-code", "Actie totaaloverzicht", *POLICY_COLUMNS}
        missing = required.difference(reader.fieldnames or [])
        if missing:
            raise ValueError(f"{reference.name}: missing reference columns {sorted(missing)}")
        for row in reader:
            name = text(row["Bestuur"])
            nis = integer(row["NIS-code"], f"{reference.name}: NIS-code")
            if len(nis) != 5 or not name:
                raise ValueError(f"{reference.name}: invalid municipality {name!r}, NIS {nis!r}")
            if name in names and names[name] != nis:
                raise ValueError(f"{reference.name}: conflicting NIS codes for {name}")
            names[name] = nis
            action = extract_code_from_text(row["Actie totaaloverzicht"])
            if action and text(row["Beleidsdomein"]):
                policy = tuple(text(row[column]) for column in POLICY_COLUMNS)
                policies.setdefault((nis, action), set()).add(policy)
    if not names:
        raise ValueError(f"{reference.name}: no municipalities found")
    return names, policies


def convert_workbook(
    workbook: Path, output: Path, reference: Path, overwrite: bool = False
) -> dict[str, int]:
    review_path = output.with_name(f"{output.stem}_unclassified.csv")
    if output.suffix.lower() != ".csv":
        raise ValueError("Output must have a .csv extension")
    for target in (output, review_path):
        if target.resolve() in {workbook.resolve(), reference.resolve()}:
            raise ValueError(f"Refusing to overwrite an input file: {target}")
        if target.exists() and not overwrite:
            raise FileExistsError(f"{target} already exists; use --overwrite to replace it")

    names, policies = load_reference(reference)
    converted: list[dict[str, str]] = []
    review: dict[tuple[str, str], dict[str, str]] = {}
    seen: set[tuple[str, str, str]] = set()
    matched = 0
    wb = load_workbook(workbook, read_only=True, data_only=True)
    try:
        sheet = wb.worksheets[0]
        source_rows = enumerate(sheet.iter_rows(values_only=True), start=1)
        headers: list[str] | None = None
        for _, values in source_rows:
            candidate = [text(value) for value in values]
            if set(REQUIRED_COLUMNS).issubset(candidate):
                headers = candidate
                break
        if headers is None:
            raise ValueError(f"{workbook.name}: no header with columns {REQUIRED_COLUMNS}")
        if len(set(headers)) != len(headers):
            raise ValueError(f"{workbook.name}: duplicate column headers")

        for row_number, values in source_rows:
            if all(value is None for value in values):
                continue
            source = dict(zip(headers, values))
            label = f"{workbook.name}, row {row_number}"
            name = text(source["Bestuur"])
            if name not in names:
                raise ValueError(f"{label}: unknown municipality {name!r}")
            nis = names[name]
            if "NIS-code" in source:
                supplied_nis = integer(source["NIS-code"], f"{label}: NIS-code")
                if supplied_nis != nis:
                    raise ValueError(f"{label}: NIS code {supplied_nis} disagrees with {name} ({nis})")
            if text(source["Type rapport"]) != "Meerjarenplan":
                raise ValueError(f"{label}: expected Type rapport Meerjarenplan")
            report_year = integer(source["Rapportjaar"], f"{label}: Rapportjaar")
            if report_year != "2026":
                raise ValueError(f"{label}: expected Rapportjaar 2026 for this blog")
            fiscal_year = integer(source["Boekjaar"], f"{label}: Boekjaar")
            action = extract_code_from_text(text(source["Actie totaaloverzicht"]))
            if not action:
                raise ValueError(f"{label}: missing action code")
            if "Actie code" in source and text(source["Actie code"]) != action:
                raise ValueError(f"{label}: Actie code disagrees with Actie totaaloverzicht")
            key = (nis, action, fiscal_year)
            if key in seen:
                raise ValueError(f"{label}: duplicate municipality/action/year {key}")
            seen.add(key)

            amount = numeric(source["Uitgave"], f"{label}: Uitgave")
            row = {column: "" for column in CSV_OUTPUT_COLUMNS}
            row.update({column: text(source[column]) for column in TEXT_COLUMNS})
            row.update({
                "NIS-code": nis,
                "Bestuur": name,
                "Type rapport": "Meerjarenplan",
                "Rapportjaar": report_year,
                "Boekjaar": fiscal_year,
                "Uitgave": format(amount, "f").replace(".", ","),
            })
            candidates = policies.get((nis, action), set())
            if len(candidates) == 1:
                row.update(dict(zip(POLICY_COLUMNS, next(iter(candidates)))))
                matched += 1
            else:
                # Action codes are local to a municipality, never global identifiers.
                review[(nis, action)] = {
                    "NIS-code": nis,
                    "Bestuur": name,
                    "Actie code": action,
                    "Reden": "Conflicting reference classifications" if candidates
                    else "No reference classification for this municipality/action",
                }
            converted.append(row)
    finally:
        wb.close()
    if not converted:
        raise ValueError(f"{workbook.name}: no data rows found")

    output.parent.mkdir(parents=True, exist_ok=True)
    for target, columns, rows in (
        (output, CSV_OUTPUT_COLUMNS, converted),
        (review_path, REVIEW_COLUMNS, list(review.values())),
    ):
        with target.open("w" if overwrite else "x", encoding="utf-8", newline="") as handle:
            writer = csv.DictWriter(handle, fieldnames=columns, delimiter=";")
            writer.writeheader()
            writer.writerows(rows)

    stats = {
        "rows": len(converted),
        "municipalities": len({row["NIS-code"] for row in converted}),
        "classified_rows": matched,
        "unclassified_rows": len(converted) - matched,
        "actions_to_review": len(review),
    }
    print(f"CSV written: {output}")
    print(f"Rows: {stats['rows']}; municipalities: {stats['municipalities']}")
    print(f"Existing policy classification retained for {matched} rows.")
    if review:
        print(
            f"WARNING: {stats['unclassified_rows']} rows / {len(review)} actions have no "
            f"unambiguous policy classification. They remain Overige in the app. Review: {review_path}"
        )
    return stats


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "workbook", nargs="?", type=Path, default=DATA_DIR / "update september 2026.xlsx"
    )
    parser.add_argument("--reference", type=Path, default=DATA_DIR / "data.csv")
    parser.add_argument("--output", type=Path, help="Default: workbook path with .csv extension")
    parser.add_argument("--overwrite", action="store_true", help="Replace existing output files")
    args = parser.parse_args()
    try:
        convert_workbook(
            args.workbook, args.output or args.workbook.with_suffix(".csv"),
            args.reference, args.overwrite,
        )
    except (ValueError, OSError) as error:
        parser.exit(1, f"Conversion failed: {error}\n")


if __name__ == "__main__":
    main()
