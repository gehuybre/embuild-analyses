"""Build October comparisons from complete economic-account exports, never the unfilterable BV exports."""

import argparse
import hashlib
import json
from pathlib import Path

import numpy as np
import pandas as pd
from openpyxl import load_workbook

APP = Path(__file__).resolve().parents[1]
SOURCE_DIR = APP.parent / "gemeentelijke-investeringen/data"
OUTPUT = APP / "public/data"
SOURCE_URL = "https://lokaalbestuur.vlaanderen.be/financien/bbc-dr/gemeente-financien"
SOURCES = {
    "latest": ("MJP meest recente oktober 26  - alle.xlsx", "Meerjarenplan"),
    "original": ("MJP origineel beginkrediet - alle.xlsx", "Meerjarenplan"),
    "actual": ("JR oktober 26.xlsx", "Jaarrekening"),
}
EXCLUDED_BV = (
    "MJP meest recente oktober 26 - alle - BV.xlsx",
    "MJP origineel - alle - BV.xlsx",
)
SCOPE = {
    "I.1.B.1.a": "Terreinen en gebouwen",
    "I.1.B.1.b": "Wegen en overige infrastructuur",
    "I.1.B.1.e": "Erfgoed",
    "I.1.B.2.a": "Onroerende goederen",
    "I.1.D": "Toegestane investeringssubsidies",
}
REQUIRED = (
    "NIS-code", "Type rapport", "Rapportjaar", "Boekjaar",
    "Niveau 2", "Niveau 5", "Alg. rekening", "Uitgave",
)


def fingerprint(path):
    return {"filename": path.name, "sha256": hashlib.sha256(path.read_bytes()).hexdigest()}


def inspect_bv_workbook(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        sheet.reset_dimensions()
        preamble = []
        iterator = sheet.iter_rows(values_only=True)
        for row in iterator:
            if row and row[0] == "NIS-code":
                return {
                    **fingerprint(path),
                    "columns": list(row),
                    "export_warnings": [text for text in preamble if "limited to" in text.lower() or "omitted" in text.lower()],
                    "reason": "No economic-account classification; not usable for construction filtering",
                }
            preamble.extend(str(value) for value in row if value is not None)
    finally:
        iterator.close()
        workbook.close()
    raise ValueError(f"{path.name}: NIS-code header missing")


def read_workbook(path):
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        sheet = workbook.active
        # BBC exports incorrectly declare A1:A1 as their used range.
        sheet.reset_dimensions()
        iterator = sheet.iter_rows(values_only=True)
        preamble = []
        for row in iterator:
            if row and row[0] == "NIS-code":
                columns = list(row)
                break
            preamble.extend(str(value) for value in row if value is not None)
        else:
            raise ValueError(f"{path.name}: NIS-code header missing")
        warnings = [text for text in preamble if "limited to" in text.lower() or "omitted" in text.lower()]
        if warnings:
            raise ValueError(f"{path.name}: truncated export: {'; '.join(warnings)}")
        frame = pd.DataFrame(iterator, columns=columns)
    finally:
        iterator.close()
        workbook.close()
    return frame, "\n".join(preamble)


def filter_construction(frame, report_type):
    missing = set(REQUIRED) - set(frame.columns)
    if missing:
        raise ValueError(f"Missing economic-account columns: {sorted(missing)}")
    data = frame[list(REQUIRED)].copy()
    data["NIS-code"] = data["NIS-code"].astype(str).str.strip()
    if not data["NIS-code"].str.fullmatch(r"\d{5}").all():
        raise ValueError("Invalid municipality codes")
    if set(data["Type rapport"]) != {report_type}:
        raise ValueError(f"Expected report type {report_type}")
    for column in ("Rapportjaar", "Boekjaar", "Uitgave"):
        data[column] = pd.to_numeric(data[column], errors="raise")
        if not np.isfinite(data[column]).all():
            raise ValueError(f"Non-finite or missing {column}")
    for column in ("Rapportjaar", "Boekjaar"):
        if not (data[column] % 1 == 0).all():
            raise ValueError(f"Non-integer {column}")
        data[column] = data[column].astype(int)
    keys = ["NIS-code", "Rapportjaar", "Boekjaar", "Alg. rekening"]
    if data.duplicated(keys).any():
        raise ValueError("Duplicate municipality/plan/year/account cells")
    if set(data["Niveau 2"]) != {"I.1 Investeringsuitgaven"}:
        raise ValueError("Unexpected expenditure level")
    if data["Niveau 5"].isna().any():
        raise ValueError("Missing construction classification")
    data = data[data["NIS-code"].str[0].isin(list("12347"))].copy()
    if report_type == "Meerjarenplan":
        if set(data.Rapportjaar) != {2014, 2020, 2026}:
            raise ValueError("Expected the three plan years 2014, 2020 and 2026")
        data = data[data.Boekjaar.between(data.Rapportjaar, data.Rapportjaar + 5)].copy()
        expected_years = set(range(2014, 2032))
    else:
        if not (data.Rapportjaar == data.Boekjaar).all():
            raise ValueError("Annual account report year must equal book year")
        expected_years = set(range(2014, 2026))
    if set(data.Boekjaar) != expected_years:
        raise ValueError(f"Unexpected or missing book years for {report_type}")
    coverage = data.groupby("Boekjaar")["NIS-code"].nunique().to_dict()
    level_codes = data["Niveau 5"].str.split().str[0]
    selected = data[level_codes.isin(SCOPE)].copy()
    # Preserve the previous REK pipeline's positive-account-cell convention.
    positive = selected[selected.Uitgave > 0].copy()
    totals = positive.groupby("Boekjaar").Uitgave.sum().reindex(sorted(expected_years), fill_value=0)
    audit = {
        "within_plan_rows": len(data),
        "construction_rows": len(selected),
        "excluded_non_positive_rows": int((selected.Uitgave <= 0).sum()),
        "excluded_non_positive_amount": float(selected.loc[selected.Uitgave <= 0, "Uitgave"].sum()),
        "positive_construction_total": float(positive.Uitgave.sum()),
    }
    return totals.to_dict(), coverage, positive, audit


def build(source_dir, output):
    output.mkdir(parents=True, exist_ok=True)
    totals, coverages, source_audit, cell_data = {}, {}, {}, {}
    for key, (filename, report_type) in SOURCES.items():
        path = source_dir / filename
        input_fingerprint = fingerprint(path)
        frame, filters = read_workbook(path)
        if key == "original" and "Origineel-Beginkrediet" not in filters:
            raise ValueError("Original export is missing its original/beginning-credit source filter")
        if key != "original" and "[Meest recente] is Ja" not in filters:
            raise ValueError(f"{filename}: most-recent source filter missing")
        totals[key], coverages[key], positive, audit = filter_construction(frame, report_type)
        cell_data[key] = positive
        source_audit[key] = {
            **input_fingerprint, "raw_rows": len(frame), "source_filters": filters, **audit,
        }
    rows = []
    for year in range(2014, 2032):
        rows.append({
            "Boekjaar": year,
            "Rapportjaar": 2014 + ((year - 2014) // 6) * 6,
            **{key: totals[key].get(year) for key in SOURCES},
            "municipalities": {key: coverages[key].get(year) for key in SOURCES},
        })
    warnings = [
        "De BV-exports zijn niet gebruikt: economische rekeningcodes ontbreken, zodat de bouwscope niet kan worden gefilterd. "
        "De recente BV-export is bovendien afgekapt op 150.000 rijen.",
        "De meest recente plannen kunnen voor verstreken jaren al zijn aangepast aan de jaarrekening. "
        "De vergelijking is geen meting van de uitvoering van het oorspronkelijke plan.",
        "Waarden van nul of lager per economische rekening worden weggelaten, conform de eerdere REK-verwerking. "
        "Dit is dus geen netto saldo inclusief negatieve boekingen.",
    ]
    for row in rows:
        counts = row["municipalities"]
        present = {key: count for key, count in counts.items() if count is not None}
        if len(set(present.values())) > 1 or any(count < 285 for count in present.values()):
            detail = ", ".join(f"{key}: {count}" for key, count in present.items())
            warnings.append(
                f"Boekjaar {row['Boekjaar']}: verschillende of onvolledige gemeentelijke dekking ({detail}). "
                "Totalen zijn niet gecorrigeerd voor ontbrekende besturen."
            )
    excluded = [inspect_bv_workbook(source_dir / filename) for filename in EXCLUDED_BV]
    if not excluded[0]["export_warnings"]:
        raise ValueError("The recent BV export has changed: review completeness and filtering before updating the warning")
    dataset = {
        "version": "oktober 2026",
        "source_url": SOURCE_URL,
        "unit": "euro",
        "price_basis": "Bedragen uit de bron, zonder bijkomende inflatiecorrectie",
        "scope": SCOPE,
        "warnings": warnings,
        "rows": rows,
    }
    snapshot = {
        "source_url": SOURCE_URL,
        "acquisition": "Door de gebruiker aangeleverde BBC-DR exports, oktober 2026",
        "retrieval_time": None,
        "licence": "Niet vermeld in de aangeleverde exports",
        "sources": source_audit,
        "excluded_sources": excluded,
        "formula": "Som van positieve Uitgave per economische rekening na Niveau 5-bouwfilter; boekjaar binnen eigen zesjarig plan",
        "scope": SCOPE,
        "warnings": warnings,
        "claims": [
            {
                "id": f"construction.{key}.{row['Boekjaar']}",
                "value": row[key],
                "unit": "EUR",
                "book_year": row["Boekjaar"],
                "plan_year": row["Rapportjaar"] if key != "actual" else None,
                "status": "reported_actual" if key == "actual" else "planned",
                "source_id": key,
                "source_file_sha256": source_audit[key]["sha256"],
                "source_rows": f"construction_{key}.csv, Boekjaar={row['Boekjaar']}",
                "municipalities": row["municipalities"][key],
            }
            for row in rows for key in SOURCES if row[key] is not None
        ],
    }
    for key in SOURCES:
        if fingerprint(source_dir / SOURCES[key][0])["sha256"] != source_audit[key]["sha256"]:
            raise ValueError("Source workbook changed during processing")
    for key, positive in cell_data.items():
        positive.to_csv(output / f"construction_{key}.csv", index=False)
    for filename, content in (("yearly_comparison.json", dataset), ("source_snapshot.json", snapshot)):
        (output / filename).write_text(json.dumps(content, ensure_ascii=False, indent=2, allow_nan=False) + "\n", encoding="utf-8")
    print(f"Generated {len(rows)} years from three filtered economic-account exports in {output}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--source-dir", type=Path, default=SOURCE_DIR)
    parser.add_argument("--output", type=Path, default=OUTPUT)
    args = parser.parse_args()
    build(args.source_dir, args.output)
