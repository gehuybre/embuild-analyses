#!/usr/bin/env python3
"""ETL sociaal-wonen: VMSW-planning (KTP en MJP, FS4-bladen) naar public/data/*.json.

Privacyregel (structureel): projectdetails bestaan enkel voor de KT-horizon. Voor LT wordt nooit een
projectveld weggeschreven. Zie PROJECT_LEVEL_HORIZONS en scripts/validate_public_data.py.

Aanpak:
- Leidende bron: de xlsx-bestanden in data/ (de CSV-exports daar worden genegeerd).
- Enkel de FS4-bladen (Nieuwbouw, Renovatie). De SSI-bladen overlappen met FS4 en worden niet meegeteld.
- Alles wat niet eenduidig te parsen of te koppelen is, laat het script falen of komt in het kwaliteitsrapport.
"""
from __future__ import annotations

import csv
import json
import re
import subprocess
import sys
import time
from collections import Counter, defaultdict
from datetime import date, datetime
from pathlib import Path

import openpyxl

from geo_lib import (APP_DIR, SvgProjector, build_current_municipalities, name_index, norm,
                     province_of)
from shapely.ops import unary_union

DATA = APP_DIR / "data"
REF = APP_DIR / "reference"
OUT = APP_DIR / "public/data"

# Bronbestanden worden automatisch gevonden: de nieuwste `JJJJ_MM_DD_KTP.xlsx` en `JJJJ_MM_DD_MJP.xlsx` in data/.
HORIZON_META = {
    "kt": {"pattern": "*_KTP.xlsx", "label": "Korte termijn", "omschrijving": "Verrichtingen die klaar zijn voor aanbesteding"},
    "lt": {"pattern": "*_MJP.xlsx", "label": "Lange termijn", "omschrijving": "Verrichtingen waarvoor een voorontwerp werd geadviseerd"},
}
HORIZONS: dict[str, dict] = {}


def find_sources() -> None:
    """Zoekt per horizon de nieuwste xlsx (op bestandsnaam) en leidt de peildatum af uit het datumvoorvoegsel."""
    for hz, m in HORIZON_META.items():
        files = sorted(p for p in DATA.glob(m["pattern"]) if not p.name.startswith("~$"))
        if not files:
            raise SystemExit(f"Geen bronbestand {m['pattern']} gevonden in {DATA}")
        f = files[-1]
        d = re.match(r"(\d{4})_(\d{2})_(\d{2})_", f.name)
        if not d:
            raise SystemExit(f"{f.name}: bestandsnaam moet beginnen met JJJJ_MM_DD_ (bv. 2026_08_14_{m['pattern'][2:]})")
        HORIZONS[hz] = {"label": m["label"], "omschrijving": m["omschrijving"], "file": f.name,
                        "peildatum": f"{d.group(1)}-{d.group(2)}-{d.group(3)}"}


def peildatum() -> str:
    return max(h["peildatum"] for h in HORIZONS.values())


# Enkel deze horizons mogen projectdetails publiceren (beslissing Gert, 2026-10-01).
PROJECT_LEVEL_HORIZONS = {"kt"}

SHEETS = {"FS4 Nieuwbouw": "nieuwbouw", "FS4 Renovatie": "renovatie"}
HEADER = ["Woonproject", "Verrichting", "Initiatiefnemers", None, "Provincie", "Gemeente", "Projectomschrijving",
          "Datum Beslissing", "Huur", "Extra info", "Type", "Kostprijs", "Maximumprijs VMSW", "Bedrag UP"]
TYPES = {
    "NB": "Nieuwbouw",
    "VVP": "Vervangingsbouw verworven patrimonium",
    "RVP": "Renovatie verworven patrimonium",
    "GW": "Aankoop goede woning",
    "VEP": "Vervangingsbouw eigen patrimonium",
    "REP": "Renovatie eigen patrimonium",
    "ONBEKEND": "Type niet vermeld in de bron",
}
PROVINCE_FIX = {"LIMBUG": "LIMBURG"}
NO_MUNI = {"geen": "Gemeente niet gespecificeerd", "meerdere": "Meerdere gemeenten"}

report: list[str] = []  # regels voor het kwaliteitsrapport


def num(v, where: str) -> float | None:
    if v is None or v == "":
        return None
    if isinstance(v, (int, float)):
        return float(v)
    s = str(v).replace("€", "").replace("\xa0", "").replace(" ", "").strip()
    if s in ("-", ""):
        return None
    if re.fullmatch(r"-?[\d.]+,\d+|-?[\d.]+", s):
        return float(s.replace(".", "").replace(",", "."))
    raise SystemExit(f"Onleesbaar getal {v!r} bij {where}")


def read_csv(path: Path) -> list[dict]:
    with open(path, newline="", encoding="utf-8") as f:
        return list(csv.DictReader(f))


def strip_prefix(n: str) -> str:
    """'woonmaatschappijdenderzuid' en 'wmdenderzuid' -> 'denderzuid' (n is al genormaliseerd)."""
    for pre in ("woonmaatschappij", "wm"):
        if n.startswith(pre) and len(n) > len(pre) + 2:
            return n[len(pre):]
    return n


AUTO_MATCHED: set[str] = set()


def load_reference():
    legend = {l["id"]: l for l in json.loads((REF / "pdf_extract/legend.json").read_text(encoding="utf-8"))}
    aliases = {norm(r["alias"]): r["wm_id"] for r in read_csv(REF / "wm_aliases.csv")}
    for wid, l in legend.items():  # naamvarianten van de legende zelf worden automatisch herkend
        for k in (norm(l["naam"]), strip_prefix(norm(l["naam"]))):
            aliases.setdefault(k, wid)
    unknown_ids = set(aliases.values()) - set(legend)
    if unknown_ids:
        raise SystemExit(f"wm_aliases.csv verwijst naar onbekende id's: {unknown_ids}")
    return legend, aliases, read_csv(REF / "gemeente_fixes.csv"), read_csv(REF / "gemeente_wm_overrides.csv")


def wm_for(name: str, aliases: dict, where: str) -> str:
    n = norm(name or "")
    if n in aliases:
        return aliases[n]
    if strip_prefix(n) in aliases:  # bv. "Woonmaatschappij X" in de bron en "WM X" in de legende
        AUTO_MATCHED.add(f"{name} -> {aliases[strip_prefix(n)]}")
        return aliases[strip_prefix(n)]
    # "Wonen in Limburg - Kantoor X (Koop)" en dergelijke
    for a, wid in aliases.items():
        if a and n.startswith(a):
            return wid
    raise SystemExit(f"Onbekende initiatiefnemer {name!r} ({where}); voeg toe aan reference/wm_aliases.csv")


def read_horizon(hz: str, legend, aliases, fixes, idx):
    wb = openpyxl.load_workbook(DATA / HORIZONS[hz]["file"], data_only=True)
    records, totals = [], defaultdict(lambda: defaultdict(float))
    text_adj = defaultdict(lambda: defaultdict(float))  # tekstcellen die de SUM van het blad overslaat
    for sheet, group in SHEETS.items():
        ws = wb[sheet]
        header = [c.value for c in ws[4]][:14]
        if header != HEADER:
            raise SystemExit(f"{hz}/{sheet}: onverwachte kolomkoppen {header}")
        for rown, r in enumerate(ws.iter_rows(min_row=5, max_col=14, values_only=True), start=5):
            where = f"{hz}/{sheet}/rij {rown}"
            if not r[0] and not r[2]:
                # totaalrij (formule): bewaren voor de controle
                if r[8] is not None:
                    for k, c in (("huur", 8), ("kostprijs", 11), ("maximumprijs_vmsw", 12), ("bedrag_up", 13)):
                        v = num(r[c], where)
                        if v is not None:  # niet elk veld heeft een som in de totaalrij
                            totals[group][k] += v
                continue
            wm = wm_for(r[2], aliases, where)
            for k, c in (("huur", 8), ("kostprijs", 11), ("maximumprijs_vmsw", 12), ("bedrag_up", 13)):
                if isinstance(r[c], str) and num(r[c], where) is not None:
                    text_adj[group][k] += num(r[c], where)
                    report.append(f"- {where}: {k} staat als tekst in de bron ({r[c]!r}). De totaalrij van het blad telt die cel niet mee; wij wel.")
            gem_raw = r[5].strip() if isinstance(r[5], str) else None
            omschr = (r[6] or "").strip()

            # expliciete correcties (reference/gemeente_fixes.csv)
            fixed_reason = None
            for fx in fixes:
                if fx["initiatiefnemer"] and norm(fx["initiatiefnemer"]) in norm(r[2] or "") \
                        and fx["omschrijving_bevat"].lower() in omschr.lower():
                    gem_raw, fixed_reason = fx["gemeente_correct"], fx["reden"]
            nis, status = None, "ok"
            if gem_raw is None:
                status = "geen"
            elif "/" in gem_raw:
                status = "meerdere"
            elif norm(gem_raw) in idx and len(idx[norm(gem_raw)]) == 1:
                nis = idx[norm(gem_raw)][0]["nis"]
            else:
                raise SystemExit(f"Gemeente {gem_raw!r} niet gevonden of dubbelzinnig ({where})")

            prov_src = PROVINCE_FIX.get((r[4] or "").strip(), (r[4] or "").strip())
            prov = province_of(nis) if nis else prov_src
            if nis and prov_src != prov and not fixed_reason:
                report.append(f"- {where}: provincie in bron ({prov_src}) wijkt af van die van {gem_raw} ({prov}); provincie afgeleid uit de gemeente.")
            if fixed_reason:
                report.append(f"- {where}: gemeente gecorrigeerd naar {gem_raw}. {fixed_reason}")

            typ = (r[10] or "").strip() or "ONBEKEND"
            if typ not in TYPES:
                raise SystemExit(f"Onbekend type {typ!r} ({where})")
            d = r[7]
            records.append({
                "horizon": hz, "groep": group, "wm": wm, "nis": nis, "status": status,
                "gemeente_label": NO_MUNI.get(status) if status != "ok" else None,
                "provincie": prov, "type": typ,
                "huur": num(r[8], where), "kostprijs": num(r[11], where),
                "maximumprijs_vmsw": num(r[12], where), "bedrag_up": num(r[13], where),
                "datum": d.date().isoformat() if isinstance(d, datetime) else None,
                "omschrijving": omschr, "procedure": (r[9] or "").strip() or None if isinstance(r[9], str) else None,
                "initiatiefnemer_bron": (r[2] or "").strip(),
            })
            if typ == "ONBEKEND":
                report.append(f"- {where}: geen type in de bron ({r[2]}); getoond als 'Type niet vermeld'.")
            if status != "ok":
                report.append(f"- {where}: {NO_MUNI[status].lower()} ({r[2]}, bron: {r[5]!r}); telt mee in totalen en tabel, niet op de kaart.")
            if "koop" in (r[2] or "").lower():
                report.append(f"- {where}: initiatiefnemer '{r[2]}' vermeldt 'Koop'; Huur = {r[8]}. Telt mee zoals in de bron.")
    # verwacht totaal = bladtotaal + tekstcellen die het blad zelf niet meetelt
    expected = {g: {k: v + text_adj[g].get(k, 0.0) for k, v in t.items()} for g, t in totals.items()}
    return records, expected


def aggregate(records):
    agg = defaultdict(lambda: {"huur": 0.0, "kostprijs": 0.0, "maximumprijs_vmsw": 0.0, "bedrag_up": 0.0,
                               "n": 0, "n_maximumprijs": 0, "jaren": []})
    for r in records:
        key = (r["wm"], r["nis"], r["status"], r["provincie"], r["groep"], r["type"], r["procedure"] or "")
        a = agg[key]
        a["huur"] += r["huur"] or 0
        a["kostprijs"] += r["kostprijs"] or 0
        a["bedrag_up"] += r["bedrag_up"] or 0
        if r["maximumprijs_vmsw"]:
            a["maximumprijs_vmsw"] += r["maximumprijs_vmsw"]
            a["n_maximumprijs"] += 1
        a["n"] += 1
        if r["datum"]:
            a["jaren"].append(int(r["datum"][:4]))
    rows = []
    for (wm, nis, status, prov, groep, typ, proc), a in sorted(agg.items(), key=lambda kv: tuple(str(x) for x in kv[0])):
        rows.append({
            "wm": wm, "nis": nis, "gemeente_label": NO_MUNI.get(status) if status != "ok" else None,
            "provincie": prov, "groep": groep, "type": typ, "procedure": proc or None,
            "huur": round(a["huur"]), "kostprijs": round(a["kostprijs"]),
            "maximumprijs_vmsw": round(a["maximumprijs_vmsw"]), "bedrag_up": round(a["bedrag_up"]),
            "n": a["n"], "n_maximumprijs": a["n_maximumprijs"],
            "jaar_min": min(a["jaren"]) if a["jaren"] else None, "jaar_max": max(a["jaren"]) if a["jaren"] else None,
        })
    return rows


def check_totals(hz, records, totals):
    """Totalen uit de bladen (totaalrij, plus tekstcellen die het blad niet meetelt) moeten kloppen met onze som."""
    lines = []
    for group, t in totals.items():
        for k, expected in t.items():
            got = sum((r[k] or 0) for r in records if r["groep"] == group)
            ok = abs(got - expected) < 1.0
            lines.append((hz, group, k, expected, got, ok))
            if not ok:
                raise SystemExit(f"Totaal {hz}/{group}/{k}: blad {expected} <> ingelezen {got}")
    return lines


def resolve_gemeente_wm(munis, legend, overrides, labels, data_by_nis):
    """Werkingsgebied per gemeente: override > PDF-kleur > planningsdata. Geeft ook de reviewlijst."""
    ov = {o["nis"]: o for o in overrides}
    pdf = defaultdict(list)
    for r in labels:
        pdf[r["nis"]].append(r)
    out, review = [], []
    for m in munis:
        nis = m["nis"]
        data_wms = data_by_nis.get(nis, {})
        cand = pdf.get(nis, [])
        pdf_wms = {c["wm_id"] for c in cand}
        worst = max((float(c["afstand"]) for c in cand), default=None)
        reasons, source, wm = [], None, None
        if nis in ov:
            wm, source = ov[nis]["wm_id"], "override"
            reasons.append(f"override: {ov[nis]['reden']}")
        elif len(pdf_wms) == 1:
            wm, source = next(iter(pdf_wms)), "pdf"
            if worst is not None and worst > 4:
                reasons.append(f"onzekere kleurmatch (afstand {worst}, tweede keuze {cand[0]['tweede_wm_id']})")
        elif len(pdf_wms) > 1:
            reasons.append("fusiegemeente met verschillende kleuren op de kaart: " + ", ".join(sorted(pdf_wms)))
            top = max(data_wms, key=data_wms.get) if data_wms else None
            wm, source = (top, "data") if top else (sorted(pdf_wms)[0], "pdf")
        if wm is None:
            reasons.append("geen PDF-label en geen planningsdata")
        if wm and data_wms and wm not in data_wms and source != "override":
            reasons.append("planning noemt andere woonmaatschappij(en): " + ", ".join(sorted(data_wms)))
            if source == "pdf" and worst is not None and worst > 4 and len(data_wms) == 1:
                wm, source = next(iter(data_wms)), "data"
                reasons.append("kleurmatch onzeker en planning eenduidig: planning gekozen")
        out.append({"nis": nis, "naam": m["name"], "provincie": m["province"], "wm_id": wm, "bron": source})
        if reasons and source != "override":
            review.append({"nis": nis, "gemeente": m["name"], "gekozen_wm": wm, "bron": source,
                           "pdf_wm": "|".join(sorted(pdf_wms)), "kleurafstand": worst if worst is not None else "",
                           "planning_wm": "|".join(sorted(data_wms)), "reden": "; ".join(reasons),
                           "definitief_wm": ""})
    return out, review


DOWNLOAD_NAME = "sociaal-wonen-alle-gegevens.xlsx"
PROV_LABEL = {"ANTWERPEN": "Antwerpen", "LIMBURG": "Limburg", "OOST-VLAANDEREN": "Oost-Vlaanderen",
              "VLAAMS-BRABANT": "Vlaams-Brabant", "WEST-VLAANDEREN": "West-Vlaanderen"}
VELDOMSCHRIJVING = {
    "huur": "Aantal huurwoningen dat zal gerealiseerd worden",
    "kostprijs": "Geraamde kostprijs van de werken",
    "maximumprijs_vmsw": "FS4-plafond",
    "bedrag_up": "Subsidiabel bedrag (berekend op basis van kostprijs en maximumprijs)",
    "datum_beslissing": "Datum opname programmatie",
}
DOWNLOAD_COLUMNS = ["Planning", "Niveau", "Woonmaatschappij", "Provincie", "NIS-code", "Gemeente", "Soort werken", "Type (code)",
                    "Type verrichting", "Procedure", "Projectomschrijving", "Datum opname programmatie", "Aantal verrichtingen",
                    "Huurwoningen", "Kostprijs (EUR)", "Maximumprijs VMSW (EUR)", "Bedrag UP (EUR)"]


def build_download_xlsx(all_records: dict, wms: list, gemeenten: list) -> list[str]:
    """Een xlsx met alle gegevens op een blad (tidy: een rij per waarneming) plus een blad Toelichting.

    KT staat op projectniveau, LT geaggregeerd (zie PROJECT_LEVEL_HORIZONS). Geen dossier-ID's, ook niet voor KT."""
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.table import Table, TableStyleInfo

    wm_name = {w["id"]: w["naam"] for w in wms}
    gem_name = {g["nis"]: g["naam"] for g in gemeenten}
    rows = []
    for hz, recs in all_records.items():
        label = HORIZONS[hz]["label"]
        if hz in PROJECT_LEVEL_HORIZONS:
            for r in sorted(recs, key=lambda r: (wm_name.get(r["wm"], r["wm"]), r["nis"] or "", r["omschrijving"] or "")):
                rows.append([label, "Project", wm_name.get(r["wm"], r["wm"]), PROV_LABEL.get(r["provincie"], r["provincie"]),
                             r["nis"], gem_name.get(r["nis"]) if r["nis"] else NO_MUNI.get(r["status"]),
                             r["groep"].capitalize(), r["type"], TYPES.get(r["type"], r["type"]), r["procedure"],
                             r["omschrijving"], date.fromisoformat(r["datum"]) if r["datum"] else None, 1, r["huur"], *(round(r[k], 2) if r[k] is not None else None for k in ("kostprijs", "maximumprijs_vmsw", "bedrag_up"))])
        else:
            for a in aggregate(recs):
                rows.append([label, "Geaggregeerd", wm_name.get(a["wm"], a["wm"]), PROV_LABEL.get(a["provincie"], a["provincie"]),
                             a["nis"], gem_name.get(a["nis"]) if a["nis"] else a["gemeente_label"],
                             a["groep"].capitalize(), a["type"], TYPES.get(a["type"], a["type"]), a["procedure"],
                             None, None, a["n"], a["huur"], a["kostprijs"], a["maximumprijs_vmsw"], a["bedrag_up"]])

    wb = openpyxl.Workbook()
    wb.properties.creator = "Embuild Vlaanderen"
    wb.properties.created = wb.properties.modified = datetime.fromisoformat(peildatum())  # herhaalbare metadata
    ws = wb.active
    ws.title = "Gegevens"
    ws.append(DOWNLOAD_COLUMNS)
    for r in rows:
        ws.append(r)
    last = len(rows) + 1
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3A5F")
        c.alignment = Alignment(wrap_text=True, vertical="center")
    for i, name in enumerate(DOWNLOAD_COLUMNS, start=1):
        letter = get_column_letter(i)
        width = min(60, max(len(name), *(len(str(r[i - 1] or "")) for r in rows)) + 2) if rows else len(name) + 2
        ws.column_dimensions[letter].width = width
        if name == "Datum opname programmatie":
            for cell in ws[letter][1:]:
                cell.number_format = "yyyy-mm-dd"
        elif name in ("Huurwoningen", "Aantal verrichtingen") or name.endswith("(EUR)"):
            for cell in ws[letter][1:]:
                cell.number_format = "#,##0"
    ws.freeze_panes = "A2"
    ws.row_dimensions[1].height = 32
    tab = Table(displayName="Gegevens", ref=f"A1:{get_column_letter(len(DOWNLOAD_COLUMNS))}{max(last, 2)}")
    tab.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=True)
    ws.add_table(tab)

    t = wb.create_sheet("Toelichting")
    lines = [
        ("Sociale huurplanning Vlaanderen: alle gegevens", "titel"),
        (f"Bron: VMSW, korte termijnplanning (KTP) en meerjarenplanning (MJP) sociale huur, peildatum {date.fromisoformat(peildatum()).strftime('%d.%m.%Y')}.", None),
        ("", None),
        ("Planningen", "kop"),
        *[(f"{h['label']}: {h['omschrijving']}.", None) for h in HORIZONS.values()],
        ("KT en LT zijn twee afzonderlijke planningen. Tel ze niet bij elkaar op.", None),
        ("", None),
        ("Opbouw van het blad Gegevens", "kop"),
        ("Een rij per waarneming. Kolom Planning onderscheidt korte en lange termijn, kolom Niveau zegt waarvoor een rij staat.", None),
        ("Niveau Project (enkel korte termijn): een rij per verrichting, met projectomschrijving en datum opname programmatie.", None),
        ("Niveau Geaggregeerd (lange termijn): een rij per combinatie van woonmaatschappij, gemeente, soort werken, type en procedure. Projectdetails zijn voor de lange termijn niet beschikbaar.", None),
        ("Interne dossiernummers (Woonproject, Verrichting) staan niet in dit bestand. De SSI-bladen zijn niet inbegrepen, omdat ze overlappen met de gewone planning.", None),
        ("", None),
        ("Kolommen", "kop"),
        ("Huurwoningen: " + VELDOMSCHRIJVING["huur"] + ". Bij renovatie gaat het om bestaande woningen waarvan de ingreep beperkt kan zijn (bv. raamcontracten of beperkte renovatie). Vergelijk bedragen per woning daarom enkel binnen hetzelfde soort werken.", None),
        ("Kostprijs (EUR): " + VELDOMSCHRIJVING["kostprijs"] + ".", None),
        ("Maximumprijs VMSW (EUR): " + VELDOMSCHRIJVING["maximumprijs_vmsw"] + ". Niet voor elke rij ingevuld.", None),
        ("Bedrag UP (EUR): " + VELDOMSCHRIJVING["bedrag_up"] + ".", None),
        ("Datum opname programmatie: " + VELDOMSCHRIJVING["datum_beslissing"] + " (kolom Datum Beslissing in de bron).", None),
        ("Procedure: de kolom Extra info uit de bron, bv. CBO, Modulair Wonen, Design & Build, raamcontract, beperkte renovatie. Leeg als niets vermeld is.", None),
        ("Aantal verrichtingen: aantal verrichtingen achter de rij (1 bij niveau Project).", None),
        ("Provincie en gemeente zijn afgeleid uit de NIS-code van de gemeente, niet uit de bronkolom. Rijen zonder eenduidige gemeente hebben geen NIS-code.", None),
        ("Herhaalde huurwaarden binnen een verrichting zijn niet gecorrigeerd, zodat de totalen overeenkomen met de totaalrijen van VMSW.", None),
        ("", None),
        ("Types verrichting", "kop"),
        *[(f"{k}: {v}", None) for k, v in TYPES.items() if k != "ONBEKEND"],
    ]
    t.column_dimensions["A"].width = 120
    for text, kind in lines:
        t.append([text])
        c = t.cell(row=t.max_row, column=1)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        if kind == "titel":
            c.font = Font(bold=True, size=14)
        elif kind == "kop":
            c.font = Font(bold=True, color="1F3A5F")
    OUT.mkdir(parents=True, exist_ok=True)
    wb.save(OUT / DOWNLOAD_NAME)
    old_zip = OUT / "sociaal-wonen-alle-gegevens.zip"
    if old_zip.exists():
        old_zip.unlink()  # vervangen door de xlsx
    return [f"{DOWNLOAD_NAME}: {len(rows)} rijen"]


def write_json(name: str, obj):
    OUT.mkdir(parents=True, exist_ok=True)
    (OUT / name).write_text(json.dumps(obj, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")


def maybe_extract_pdf():
    """Draait de PDF-extractie als de legende ontbreekt of de PDF nieuwer is."""
    legend = REF / "pdf_extract/legend.json"
    pdfs = sorted((DATA).glob("Woonmaatschappijen-in-kaart*.pdf"))
    if pdfs and (not legend.exists() or pdfs[-1].stat().st_mtime > legend.stat().st_mtime):
        print("Nieuwe PDF gevonden: werkingsgebieden worden opnieuw uitgelezen")
        r = subprocess.run([sys.executable, str(Path(__file__).with_name("extract_wm_pdf.py"))])
        if r.returncode:
            raise SystemExit("extract_wm_pdf.py faalde")


def main():
    find_sources()
    maybe_extract_pdf()
    for hz, h in HORIZONS.items():
        print(f"{hz.upper()}: {h['file']} (peildatum {h['peildatum']})")
    legend, aliases, fixes, overrides = load_reference()
    munis = build_current_municipalities()
    idx = name_index(munis)
    labels = read_csv(REF / "pdf_extract/labels_auto.csv")

    all_records: dict[str, list] = {}
    checks = []
    for hz in HORIZONS:
        recs, totals = read_horizon(hz, legend, aliases, fixes, idx)
        checks += check_totals(hz, recs, totals)
        all_records[hz] = recs
        print(f"{hz.upper()}: {len(recs)} rijen; huur {sum(r['huur'] or 0 for r in recs):,.0f}; "
              f"kostprijs {sum(r['kostprijs'] or 0 for r in recs):,.0f}")

    # data-WM's per gemeente (voor kruiscontrole)
    data_by_nis: dict[str, Counter] = defaultdict(Counter)
    for recs in all_records.values():
        for r in recs:
            if r["nis"]:
                data_by_nis[r["nis"]][r["wm"]] += r["huur"] or 1
    gemeenten, review = resolve_gemeente_wm(munis, legend, overrides, labels, data_by_nis)

    # Review-bestand met behoud van reeds ingevulde keuzes
    review_path = REF / "gemeente_wm_review.csv"
    previous = {r["nis"]: r for r in read_csv(review_path)} if review_path.exists() else {}
    for r in review:
        if r["nis"] in previous:
            r["definitief_wm"] = previous[r["nis"]].get("definitief_wm", "")
    with open(review_path, "w", newline="", encoding="utf-8") as f:
        w = csv.DictWriter(f, fieldnames=list(review[0].keys()))
        w.writeheader()
        w.writerows(review)
    # definitieve keuzes van Gert winnen
    final = {r["nis"]: r["definitief_wm"] for r in review if r.get("definitief_wm")}
    for g in gemeenten:
        if g["nis"] in final:
            if final[g["nis"]] not in legend:
                raise SystemExit(f"definitief_wm {final[g['nis']]!r} onbekend voor {g['naam']}")
            g["wm_id"], g["bron"] = final[g["nis"]], "review"
    missing = [g["naam"] for g in gemeenten if not g["wm_id"]]
    if missing:
        raise SystemExit(f"Gemeenten zonder woonmaatschappij: {missing}")

    # woonmaatschappijen
    by_wm = defaultdict(list)
    for g in gemeenten:
        by_wm[g["wm_id"]].append(g["nis"])
    wms = []
    for wid, l in sorted(legend.items(), key=lambda kv: norm(kv[1]["naam"])):
        r, g, b = l["kleur"]
        wms.append({
            "id": wid, "naam": l["naam"], "url": l["url"],
            "werkingsgebied_label": l["werkingsgebied_label"], "provincie": l["provincie_legende"],
            "kleur": f"#{r:02x}{g:02x}{b:02x}", "gemeenten": sorted(by_wm.get(wid, [])),
        })

    # geometrie (huidige gemeenten) en provincies als SVG-paden, uit dezelfde bron zodat randen aansluiten
    proj = SvgProjector([m["geom"] for m in munis])
    map_munis = []
    for m in munis:
        cx, cy = proj.centroid(m["geom"])
        map_munis.append({"nis": m["nis"], "naam": m["name"], "provincie": m["province"], "d": proj.path(m["geom"]), "cx": cx, "cy": cy})
    map_provs = []
    for p in sorted({m["province"] for m in munis}):
        geom = unary_union([m["geom"].buffer(0.00005) for m in munis if m["province"] == p])
        map_provs.append({"naam": p, "d": proj.path(geom, tolerance=0.002)})
    write_json("map.json", {"width": proj.width, "height": proj.height,
                            "gemeenten": map_munis, "provincies": map_provs})

    # planningsdata
    for hz, recs in all_records.items():
        write_json(f"{hz}_aggregated.json", {"horizon": hz, "peildatum": HORIZONS[hz]["peildatum"], "rijen": aggregate(recs)})
    for hz in PROJECT_LEVEL_HORIZONS:
        projects = [{
            "projectomschrijving": r["omschrijving"], "wm": r["wm"], "nis": r["nis"],
            "gemeente_label": r["gemeente_label"], "provincie": r["provincie"], "groep": r["groep"], "type": r["type"],
            "huur": r["huur"], "kostprijs": round(r["kostprijs"]) if r["kostprijs"] is not None else None,
            "maximumprijs_vmsw": round(r["maximumprijs_vmsw"]) if r["maximumprijs_vmsw"] is not None else None,
            "bedrag_up": round(r["bedrag_up"]) if r["bedrag_up"] is not None else None,
            "datum_beslissing": r["datum"], "procedure": r["procedure"],
        } for r in all_records[hz]]
        projects.sort(key=lambda p: (p["wm"], p["nis"] or "", p["projectomschrijving"]))
        write_json(f"{hz}_projects.json", {"horizon": hz, "peildatum": HORIZONS[hz]["peildatum"], "projecten": projects})

    write_json("woonmaatschappijen.json", wms)
    write_json("gemeenten.json", gemeenten)
    print("Download:", build_download_xlsx(all_records, wms, gemeenten)[0])

    control = {f"{hz}.{g}.{k}": round(sum((r[k] or 0) for r in recs if r["groep"] == g))
               for hz, recs in all_records.items() for g in SHEETS.values()
               for k in ("huur", "kostprijs", "maximumprijs_vmsw", "bedrag_up")}
    write_json("meta.json", {
        "peildatum": peildatum(),
        "lastUpdated": date.today().isoformat(),
        "bron": f"VMSW, korte termijnplanning (KTP) en meerjarenplanning (MJP) sociale huur, {date.fromisoformat(peildatum()).strftime('%d.%m.%Y')}",
        "werkingsgebieden": "Woonmaatschappijen in kaart, januari 2025 (kleurafleiding uit PDF, zie reference/)",
        "horizons": {k: {"label": v["label"], "omschrijving": v["omschrijving"], "peildatum": v["peildatum"]} for k, v in HORIZONS.items()},
        "projectdetails_horizons": sorted(PROJECT_LEVEL_HORIZONS),
        "types": TYPES,
        "veldomschrijving": VELDOMSCHRIJVING,
        "controletotalen": control,
    })

    # kwaliteitsrapport
    lines = ["# Kwaliteitsrapport sociaal-wonen", "", f"Gegenereerd door `scripts/process_data.py`. Bronbestanden: {", ".join(h["file"] for h in HORIZONS.values())}.", "",
             "## Controle tegen de totaalrijen in de bladen", "",
             "'Totaal in blad' = totaalrij uit het bronblad plus de bedragen die in de bron als tekst staan en door de SUM van het blad worden overgeslagen (zie rij-specifieke opmerkingen).", "",
             "| Horizon | Blad | Veld | Totaal in blad | Ingelezen | OK |", "|---|---|---|---:|---:|---|"]
    for hz, g, k, e, got, ok in checks:
        lines.append(f"| {hz.upper()} | {g} | {k} | {e:,.0f} | {got:,.0f} | {'ja' if ok else 'NEE'} |")
    lines += ["", "## Beslissingen en aandachtspunten", "",
              "- SSI-bladen zijn niet inbegrepen (overlap met FS4).",
              "- Huur wordt opgeteld zoals in de bron. Herhaalde huurwaarden binnen een verrichting zijn niet gecorrigeerd, zodat de totalen overeenkomen met de officiele totaalrij.",
              "- Provincie is afgeleid uit de gemeente (NIS), niet uit de bronkolom.",
              "- Projectdetails worden enkel voor KT gepubliceerd. Interne dossier-ID's (Woonproject, Verrichting) worden nooit gepubliceerd.",
              "", "## Rij-specifieke opmerkingen", ""] + (sorted(set(report)) or ["- geen"])
    (REF / "kwaliteitsrapport.md").write_text("\n".join(lines) + "\n", encoding="utf-8")

    if AUTO_MATCHED:
        print("Automatisch gekoppelde naamvarianten (controleer eventueel):", "; ".join(sorted(AUTO_MATCHED)))
    print(f"Gemeenten: {len(gemeenten)}, bronnen: {dict(Counter(g['bron'] for g in gemeenten))}")
    print(f"Review-lijst: {len(review)} gemeenten (reference/gemeente_wm_review.csv)")
    print("Totalen kloppen met de bladen:", all(c[-1] for c in checks))


def source_state() -> dict:
    return {p.name: p.stat().st_mtime for p in list(DATA.glob("*.xlsx")) + list(DATA.glob("*.pdf")) if not p.name.startswith("~$")}


def run_once() -> bool:
    try:
        main()
    except SystemExit as e:
        if e.code not in (0, None):
            print(f"\nVerwerking gestopt: {e.code}", file=sys.stderr)
            return False
    r = subprocess.run([sys.executable, str(Path(__file__).with_name("validate_public_data.py"))])
    return r.returncode == 0


if __name__ == "__main__":
    if "--watch" in sys.argv:
        # Verwerkt automatisch zodra een nieuwe of gewijzigde xlsx/pdf in data/ verschijnt.
        sys.stdout.reconfigure(line_buffering=True)
        print(f"Bewaakt {DATA} (Ctrl+C om te stoppen)")
        seen = {}
        while True:
            now = source_state()
            if now != seen:
                time.sleep(2)  # wacht tot het kopieren klaar is
                seen = source_state()
                print(f"\n[{datetime.now():%H:%M:%S}] wijziging gevonden, verwerken...")
                ok = run_once()
                print("Klaar, alles in orde." if ok else "Mislukt: zie de melding hierboven. Pas aan en sla opnieuw op.")
            time.sleep(3)
    else:
        sys.exit(0 if run_once() else 1)
