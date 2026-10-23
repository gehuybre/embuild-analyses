"""Extract Revtool locally without executing VBA, refreshing queries or recalculating."""

import argparse
import base64
import csv
import io
import json
import math
import re
import struct
import warnings
from collections import Counter
from datetime import date, datetime, time
from decimal import Decimal
from hashlib import sha256
from importlib.metadata import version
from pathlib import Path
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from oletools.olevba import VBA_Parser
from openpyxl import load_workbook
from openpyxl.utils.cell import range_boundaries

APP = Path(__file__).resolve().parents[1]
REGIME = "offerten vanaf 11/06/2007"
SERIES_SHEETS = {
    "Standaardindexen",
    "Materiaalindexen o.w.",
    "Grymafer indexen",
    "Indexen goederen & diensten",
    "Indexeringen Lonen per PC",
}
SERIES_COLUMNS = [
    "sheet", "table", "series", "date", "cell", "value", "status", "raw_value", "formula",
]
COMPARISON_COLUMNS = [
    "date", "series", "cell", "employer_size", "category", "workbook_value",
    "pdf_value", "difference", "status", "effective_date", "valid_until",
    "circular", "source_pdf", "source_page",
]


def scalar(value):
    if isinstance(value, (datetime, date, time)):
        return value.isoformat()
    if value is None or isinstance(value, (str, bool, int, float)):
        return value
    raise TypeError(f"Unsupported cell value: {type(value).__name__}")


def formula_text(cell) -> str | None:
    if cell.data_type != "f":
        return None
    if isinstance(cell.value, str):
        return cell.value
    if hasattr(cell.value, "text"):
        return cell.value.text
    raise ValueError(f"Unsupported formula at {cell.coordinate}: {cell.value!r}")


def safe_name(name: str) -> str:
    result = re.sub(r"[^A-Za-z0-9_.-]+", "_", name).strip("._")
    if not result:
        raise ValueError(f"Cannot create filename for {name!r}")
    return result


def find_latest_workbook(data_dir: Path) -> Path:
    candidates = []
    for path in data_dir.glob("revtool-nl-versie-*.xlsm"):
        match = re.fullmatch(r"revtool-nl-versie-(\d{4}-\d{2}-\d{2})\.xlsm", path.name)
        if match:
            candidates.append((date.fromisoformat(match[1]), path))
    if not candidates:
        raise FileNotFoundError(f"No dated Revtool workbook in {data_dir}; supply --input")
    return max(candidates)[1]


def write_json(path: Path, payload) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        json.dumps(payload, ensure_ascii=False, indent=2, allow_nan=False) + "\n",
        encoding="utf-8",
    )


def write_csv(path: Path, columns: list[str], rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=columns)
        writer.writeheader()
        writer.writerows(rows)


def extract_series(sheet, formulas) -> tuple[list[dict], list[dict]]:
    """Use structured tables, not display/helper ranges or inferred forward fills."""
    series = []
    catalogue = []
    if sheet.title not in SERIES_SHEETS:
        return series, catalogue
    for table in sheet.tables.values():
        left, top, right, bottom = range_boundaries(table.ref)
        for column in range(left + 1, right + 1):
            header = sheet.cell(top, column)
            if header.value == "Week nummer":
                continue
            if not isinstance(header.value, str) or not header.value.strip():
                raise ValueError(f"Missing series header in {sheet.title}!{header.coordinate}")
            observations = []
            dates = set()
            for row in range(top + 1, bottom + 1 - (table.totalsRowCount or 0)):
                period = sheet.cell(row, left).value
                if period is None:
                    continue
                if not isinstance(period, (date, datetime)):
                    raise ValueError(f"Non-date period in {sheet.title}!{sheet.cell(row, left).coordinate}")
                period_date = period.date().isoformat() if isinstance(period, datetime) else period.isoformat()
                if period_date in dates:
                    raise ValueError(f"Duplicate period in {sheet.title}/{table.name}: {period_date}")
                dates.add(period_date)
                cell = sheet.cell(row, column)
                value = cell.value
                numeric = (
                    isinstance(value, (int, float))
                    and not isinstance(value, bool)
                    and math.isfinite(value)
                )
                status = (
                    "numeric" if numeric else
                    "missing" if value is None else
                    "excel_error" if cell.data_type == "e" else "non_numeric"
                )
                observations.append({
                    "sheet": sheet.title,
                    "table": table.name,
                    "series": str(header.value),
                    "date": period_date,
                    "cell": cell.coordinate,
                    "value": value if numeric else None,
                    "status": status,
                    "raw_value": scalar(value),
                    "formula": formula_text(formulas.cell(row, column)),
                })
            numeric_rows = [row for row in observations if row["status"] == "numeric"]
            catalogue.append({
                "sheet": sheet.title,
                "table": table.name,
                "series": str(header.value),
                "header_cell": header.coordinate,
                "number_format": sheet.cell(top + 1, column).number_format,
                "first_numeric_date": min((r["date"] for r in numeric_rows), default=None),
                "last_numeric_date": max((r["date"] for r in numeric_rows), default=None),
                "status_counts": dict(Counter(r["status"] for r in observations)),
            })
            series.extend(observations)
    return series, catalogue


def compare_labor(series: list[dict], pdf_rows: list[dict]) -> list[dict]:
    groups = {}
    seen = set()
    for row in pdf_rows:
        if row["regime"] != REGIME or row["currency"] != "EUR":
            continue
        key = (row["employer_size"], row["category"])
        unique = (*key, row["effective_date"])
        if unique in seen:
            raise ValueError(f"Duplicate PDF observation: {unique}")
        seen.add(unique)
        groups.setdefault(key, []).append(row)
    for rows in groups.values():
        rows.sort(key=lambda row: row["effective_date"])

    # These are source-label aliases, not a redefinition of the legal thresholds.
    sizes = {"< 10 p": "<10 werklieden", "van 10 tot 19 p": "10-20 werklieden",
             "\u2265 20 p": ">20 werklieden"}
    results = []
    for row in series:
        if row["sheet"] != "Standaardindexen":
            continue
        match = re.fullmatch(r'index "s" - cat ([A-D]) \((.*?)\)', row["series"])
        if not match:
            continue
        category, size = match.groups()
        if size not in sizes:
            raise ValueError(f"Unknown employer-size label: {size}")
        result = {
            "date": row["date"], "series": row["series"], "cell": row["cell"],
            "employer_size": sizes[size], "category": category,
            "workbook_value": row["value"], "status": "workbook_" + row["status"],
        }
        if row["status"] == "numeric":
            candidates = [
                item for item in groups.get((sizes[size], category), [])
                if item["effective_date"] <= row["date"]
            ]
            pdf = candidates[-1] if candidates else None
            if pdf is None or (pdf["valid_until"] and row["date"] > pdf["valid_until"]):
                result["status"] = "no_pdf_period"
            else:
                difference = Decimal(str(row["value"])) - Decimal(pdf["value"])
                result.update({
                    "pdf_value": float(pdf["value"]),
                    "difference": float(difference),
                    "status": "match" if abs(difference) <= Decimal("0.0005") else "mismatch",
                    **{key: pdf[key] for key in [
                        "effective_date", "valid_until", "circular", "source_pdf", "source_page",
                    ]},
                })
        results.append(result)
    return results


def extract_code(source: Path, output: Path) -> dict:
    modules = []
    parser = VBA_Parser(str(source))
    try:
        for _, stream, name, code in parser.extract_macros():
            target = f"vba/{len(modules) + 1:02d}_{safe_name(name)}"
            path = output / target
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_text(code, encoding="utf-8")
            modules.append({
                "name": name, "stream": stream, "file": target,
                "procedures": re.findall(
                    r"^\s*(?:(?:Public|Private|Friend|Static)\s+)?(?:Sub|Function)\s+(\w+)",
                    code, flags=re.MULTILINE | re.IGNORECASE,
                ),
            })
    finally:
        parser.close()

    package_parts = []
    queries = []
    with ZipFile(source) as archive:
        for name in archive.namelist():
            if name == "xl/vbaProject.bin" or name == "xl/connections.xml" or name.startswith("customXml/"):
                content = archive.read(name)
                target = output / "package" / name
                if not target.resolve().is_relative_to((output / "package").resolve()):
                    raise ValueError(f"Unsafe workbook package path: {name}")
                target.parent.mkdir(parents=True, exist_ok=True)
                target.write_bytes(content)
                package_parts.append(name)
                if name.startswith("customXml/") and name.endswith(".xml"):
                    root = ET.fromstring(content)
                    if root.tag == "{http://schemas.microsoft.com/DataMashup}DataMashup":
                        payload = base64.b64decode(root.text or "", validate=True)
                        format_version, size = struct.unpack("<II", payload[:8])
                        if format_version != 0 or len(payload) < 8 + size:
                            raise ValueError(f"Unsupported DataMashup header: {name}")
                        with ZipFile(io.BytesIO(payload[8:8 + size])) as mashup:
                            for part in mashup.namelist():
                                if part.endswith(".m"):
                                    filename = f"powerquery/{safe_name(Path(name).stem)}_{safe_name(Path(part).name)}"
                                    target = output / filename
                                    target.parent.mkdir(parents=True, exist_ok=True)
                                    target.write_bytes(mashup.read(part))
                                    queries.append({"source": name, "part": part, "file": filename})
    return {"vba_modules": modules, "power_query": queries, "preserved_package_parts": package_parts}


def extract(source: Path, output: Path, pdf_csv: Path) -> dict:
    source = source.resolve()
    output = output.resolve()
    if output == source.parent or output in source.parents or output in pdf_csv.resolve().parents:
        raise ValueError("Output must be a dedicated extraction directory")
    source_hash = sha256(source.read_bytes()).hexdigest()
    output.mkdir(parents=True, exist_ok=True)
    with warnings.catch_warnings(record=True) as captured:
        warnings.simplefilter("always", UserWarning)
        formulas = load_workbook(source, data_only=False, keep_links=True)
        values = load_workbook(source, data_only=True, keep_links=True)
    manifest = {
        "source": source.name,
        "sha256": source_hash,
        "source_size": source.stat().st_size,
        "dependencies": {name: version(name) for name in ["openpyxl", "oletools"]},
        "mode": "static; no macros executed, no queries refreshed, no formulas recalculated",
        "licence": "Not established from the local source; extraction does not grant redistribution rights.",
        "series_status_note": (
            "numeric means a saved numeric cell, not a verified realized observation; "
            "wage tables can contain forward projections. Missing values and errors are not filled."
        ),
        "reader_warnings": sorted({str(item.message) for item in captured}),
        "defined_names": [dict(item) for item in formulas.defined_names.values()],
        "calculation_properties": dict(formulas.calculation) if formulas.calculation else None,
        "sheets": [],
    }
    series = []
    catalogue = []
    for index, sheet in enumerate(formulas, 1):
        cached = values[sheet.title]
        stem = f"sheets/{index:02d}_{safe_name(sheet.title)}"
        cells = []
        for row in sheet:
            for cell in row:
                value = cached.cell(cell.row, cell.column)
                if cell.value is None and value.value is None:
                    continue
                formula = formula_text(cell)
                record = {
                    "cell": cell.coordinate, "value": scalar(value.value),
                    "data_type": value.data_type, "formula": formula,
                    "formula_cache_missing": formula is not None and value.value is None,
                    "number_format": cell.number_format,
                }
                if formula is not None and not isinstance(cell.value, str):
                    record["formula_attributes"] = dict(cell.value)
                if cell.hyperlink:
                    record["hyperlink"] = cell.hyperlink.target or cell.hyperlink.location
                if cell.comment:
                    record["comment"] = cell.comment.text
                cells.append(record)
        write_json(output / f"{stem}.cells.json", cells)
        csv_path = output / f"{stem}.csv"
        with csv_path.open("w", encoding="utf-8", newline="") as handle:
            writer = csv.writer(handle)
            writer.writerows([scalar(cell.value) for cell in row] for row in cached)
        manifest["sheets"].append({
            "name": sheet.title, "state": sheet.sheet_state,
            "rows": sheet.max_row, "columns": sheet.max_column,
            "csv": f"{stem}.csv", "cells": f"{stem}.cells.json",
            "populated_cells": len(cells),
            "formulas": sum(cell["formula"] is not None for cell in cells),
            "missing_formula_caches": sum(cell["formula_cache_missing"] for cell in cells),
            "excel_errors": sum(cell["data_type"] == "e" for cell in cells),
            "merged_ranges": [str(item) for item in sheet.merged_cells.ranges],
            "tables": [
                {"name": table.name, "ref": table.ref,
                 "columns": [column.name for column in table.tableColumns]}
                for table in sheet.tables.values()
            ],
        })
        observations, entries = extract_series(cached, sheet)
        series.extend(observations)
        catalogue.extend(entries)
    formulas.close()
    values.close()
    manifest.update(extract_code(source, output))
    write_csv(output / "index_series.csv", SERIES_COLUMNS, series)
    write_json(output / "series_catalogue.json", catalogue)
    with pdf_csv.open(encoding="utf-8", newline="") as handle:
        comparison = compare_labor(series, list(csv.DictReader(handle)))
    if not comparison:
        raise ValueError("No standard s categories found; check workbook sheet names, tables and labels")
    counts = dict(Counter(row["status"] for row in comparison))
    write_csv(output / "pdf_comparison.csv", COMPARISON_COLUMNS, comparison)
    write_json(output / "pdf_comparison.json", {
        "pdf_csv": pdf_csv.name, "pdf_csv_sha256": sha256(pdf_csv.read_bytes()).hexdigest(),
        "regime": REGIME, "comparison_date": "first day of workbook month; execution s, not tender S",
        "absolute_tolerance": 0.0005,
        "employer_size_note": (
            "Revtool <10 / 10-19 / >=20 mapped by source column order to existing CSV "
            "<10 / 10-20 / >20 labels; the differing threshold wording requires separate review."
        ),
        "status_counts": counts,
        "mismatches": [row for row in comparison if row["status"] == "mismatch"],
    })
    manifest["series_count"] = len(catalogue)
    manifest["series_observations"] = len(series)
    manifest["comparison_status_counts"] = counts
    if sha256(source.read_bytes()).hexdigest() != source_hash:
        raise RuntimeError("Source workbook changed during extraction")
    artifacts = [
        "index_series.csv", "series_catalogue.json", "pdf_comparison.csv", "pdf_comparison.json",
        *[sheet[key] for sheet in manifest["sheets"] for key in ["csv", "cells"]],
        *[module["file"] for module in manifest["vba_modules"]],
        *[query["file"] for query in manifest["power_query"]],
        *[f"package/{part}" for part in manifest["preserved_package_parts"]],
    ]
    manifest["artifacts"] = [
        {"file": str(path.relative_to(output)), "sha256": sha256(path.read_bytes()).hexdigest()}
        for path in sorted(output / artifact for artifact in artifacts)
    ]
    write_json(output / "manifest.json", manifest)
    return manifest


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--input", type=Path, help="Default: latest dated revtool-nl-versie-*.xlsm in data/")
    parser.add_argument("--output", type=Path, help="Default: data/revtool-extracted/<workbook-stem>/")
    parser.add_argument("--pdf-csv", type=Path, default=APP / "data/waarden-S-s.csv")
    args = parser.parse_args()
    source = args.input or find_latest_workbook(APP / "data")
    output = args.output or APP / "data/revtool-extracted" / source.stem
    manifest = extract(source, output, args.pdf_csv)
    print(f"Extracted {len(manifest['sheets'])} sheets, {len(manifest['vba_modules'])} VBA modules, "
          f"{len(manifest['power_query'])} Power Query files and {manifest['series_count']} series.")
    print(f"PDF comparison: {manifest['comparison_status_counts']}")
    for message in manifest["reader_warnings"]:
        print(f"Reader warning (original workbook unchanged): {message}")
    print(f"Output: {output.resolve()}")


if __name__ == "__main__":
    main()
