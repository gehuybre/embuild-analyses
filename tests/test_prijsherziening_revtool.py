import base64
import csv
import importlib.util
import io
import json
import struct
import tempfile
import unittest
from datetime import datetime
from hashlib import sha256
from pathlib import Path
from unittest.mock import patch
from zipfile import ZipFile

from openpyxl import Workbook
from openpyxl.worksheet.table import Table

APP = Path(__file__).resolve().parents[1] / "apps/prijsherziening-index-i-2021"
SPEC = importlib.util.spec_from_file_location("revtool", APP / "scripts/extract_revtool.py")
REVTOOL = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(REVTOOL)


def observation(value=31.074, day="2018-01-01", status="numeric", category="A", size="< 10 p"):
    return {
        "sheet": "Standaardindexen", "series": f'index "s" - cat {category} ({size})',
        "date": day, "cell": "D2", "value": value, "status": status,
    }


def pdf_row(value="31.074", **changes):
    return {
        "regime": REVTOOL.REGIME, "currency": "EUR", "employer_size": "<10 werklieden",
        "category": "A", "value": value, "effective_date": "2018-01-01",
        "valid_until": "2018-03-31", "circular": "544-test",
        "source_pdf": "/data/sources/fixture.pdf", "source_page": "2", **changes,
    }


class RevtoolTests(unittest.TestCase):
    def test_discovers_latest_version_by_filename_not_modification_time(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            latest = root / "revtool-nl-versie-2027-02-01.xlsm"
            latest.touch()
            (root / "revtool-nl-versie-2026-07-10.xlsm").touch()
            (root / "unrelated.xlsm").touch()
            self.assertEqual(REVTOOL.find_latest_workbook(root), latest)

    def test_missing_workbook_is_an_explicit_error(self):
        with tempfile.TemporaryDirectory() as tmp:
            with self.assertRaisesRegex(FileNotFoundError, "supply --input"):
                REVTOOL.find_latest_workbook(Path(tmp))

    def test_series_uses_table_bounds_and_preserves_missing_and_errors(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Standaardindexen"
        for row in [
            ["Date", "Series", "Week nummer"],
            [datetime(2026, 1, 1), 0, 1],
            [datetime(2026, 2, 1), None, 2],
            [datetime(2026, 3, 1), "#N/A", 3],
            [datetime(2026, 4, 1), "no data", 4],
        ]:
            sheet.append(row)
        sheet.add_table(Table(displayName="NewVersionTable", ref="A1:C5"))
        sheet["D1"] = "Unrelated helper"
        sheet["D2"] = 999
        rows, catalogue = REVTOOL.extract_series(sheet, sheet)
        self.assertEqual([row["status"] for row in rows],
                         ["numeric", "missing", "excel_error", "non_numeric"])
        self.assertEqual(rows[0]["value"], 0)
        self.assertIsNone(rows[2]["value"])
        self.assertEqual(rows[2]["raw_value"], "#N/A")
        self.assertEqual(len(catalogue), 1)
        self.assertEqual(catalogue[0]["last_numeric_date"], "2026-01-01")
        self.assertEqual(catalogue[0]["table"], "NewVersionTable")

    def test_non_date_table_period_is_not_silently_skipped(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Standaardindexen"
        sheet.append(["Date", "Value"])
        sheet.append(["not a date", 123])
        sheet.add_table(Table(displayName="ChangedLayout", ref="A1:B2"))
        with self.assertRaisesRegex(ValueError, "Non-date period"):
            REVTOOL.extract_series(sheet, sheet)

    def test_duplicate_series_dates_are_an_explicit_error(self):
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Standaardindexen"
        sheet.append(["Date", "Value"])
        sheet.append([datetime(2026, 1, 1), 1])
        sheet.append([datetime(2026, 1, 1), 2])
        sheet.add_table(Table(displayName="DuplicateDates", ref="A1:B3"))
        with self.assertRaisesRegex(ValueError, "Duplicate period"):
            REVTOOL.extract_series(sheet, sheet)

    def test_comparison_uses_applicability_period_and_keeps_pdf_provenance(self):
        rows = REVTOOL.compare_labor([
            observation(day="2018-02-01"),
            observation(day="2018-04-01"),
            observation(day="2017-12-01"),
            observation(value=None, status="missing"),
        ], [pdf_row()])
        self.assertEqual([r["status"] for r in rows],
                         ["match", "no_pdf_period", "no_pdf_period", "workbook_missing"])
        self.assertEqual(rows[0]["source_page"], "2")
        self.assertEqual(rows[0]["effective_date"], "2018-01-01")

    def test_latest_effective_period_wins(self):
        rows = REVTOOL.compare_labor(
            [observation(value=32, day="2018-02-01")],
            [pdf_row(), pdf_row(value="32", effective_date="2018-02-01")],
        )
        self.assertEqual(rows[0]["status"], "match")
        self.assertEqual(rows[0]["pdf_value"], 32)

    def test_comparison_threshold_is_half_last_published_decimal(self):
        rows = REVTOOL.compare_labor(
            [observation(31.0745), observation(31.074501)], [pdf_row()],
        )
        self.assertEqual([r["status"] for r in rows], ["match", "mismatch"])

    def test_comparison_covers_all_categories_and_three_sizes(self):
        sizes = [("< 10 p", "<10 werklieden"), ("van 10 tot 19 p", "10-20 werklieden"),
                 ("\u2265 20 p", ">20 werklieden")]
        observations = [
            observation(category=category, size=label) for category in "ABCD" for label, _ in sizes
        ]
        pdf = [
            pdf_row(category=category, employer_size=size) for category in "ABCD" for _, size in sizes
        ]
        result = REVTOOL.compare_labor(observations, pdf)
        self.assertEqual(len(result), 12)
        self.assertTrue(all(r["status"] == "match" for r in result))

    def test_comparison_rejects_duplicates_and_excludes_other_regimes(self):
        with self.assertRaisesRegex(ValueError, "Duplicate PDF"):
            REVTOOL.compare_labor([observation()], [pdf_row(), pdf_row()])
        result = REVTOOL.compare_labor(
            [observation()], [pdf_row(regime="older regime"), pdf_row(currency="BEF")],
        )
        self.assertEqual(result[0]["status"], "no_pdf_period")

    def test_static_vba_and_power_query_extraction(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "fixture.xlsm"
            package = io.BytesIO()
            with ZipFile(package, "w") as archive:
                archive.writestr("Formulas/Section1.m", 'section Section1; shared Test = 42;')
            payload = struct.pack("<II", 0, len(package.getvalue())) + package.getvalue()
            xml = (
                '<DataMashup xmlns="http://schemas.microsoft.com/DataMashup">'
                + base64.b64encode(payload).decode("ascii") + "</DataMashup>"
            )
            with ZipFile(source, "w") as archive:
                archive.writestr("xl/vbaProject.bin", b"raw VBA bytes")
                archive.writestr("xl/connections.xml", "<connections/>")
                archive.writestr("customXml/item1.xml", xml)
            with patch.object(REVTOOL, "VBA_Parser") as parser:
                parser.return_value.extract_macros.return_value = [
                    ("fixture", "VBA/Module1", "../Module1.bas", "Sub Example()\nEnd Sub\n"),
                ]
                result = REVTOOL.extract_code(source, root / "output")
                parser.return_value.close.assert_called_once()
            self.assertEqual(result["vba_modules"][0]["procedures"], ["Example"])
            for module in result["vba_modules"]:
                self.assertTrue((root / "output" / module["file"]).is_relative_to(root / "output/vba"))
            query = root / "output" / result["power_query"][0]["file"]
            self.assertEqual(query.read_text(), 'section Section1; shared Test = 42;')
            self.assertEqual((root / "output/package/xl/vbaProject.bin").read_bytes(), b"raw VBA bytes")

    def test_extract_preserves_hidden_sheets_formulas_and_original_bytes(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "fixture.xlsx"
            workbook = Workbook()
            workbook.active.title = "Introductie"
            sheet = workbook.create_sheet("Standaardindexen")
            sheet.sheet_state = "hidden"
            sheet.append(["Date", 'index "s" - cat A (< 10 p)', "Calculated"])
            sheet.append([datetime(2018, 1, 1), 31.074, "=B2*2"])
            sheet.add_table(Table(displayName="NewTable", ref="A1:C2"))
            workbook.save(source)
            before = source.read_bytes()
            pdf_csv = root / "pdf.csv"
            with pdf_csv.open("w", newline="") as handle:
                writer = csv.DictWriter(handle, fieldnames=list(pdf_row()))
                writer.writeheader()
                writer.writerow(pdf_row())
            output = root / "extracted"
            with patch.object(REVTOOL, "extract_code", return_value={
                "vba_modules": [], "power_query": [], "preserved_package_parts": [],
            }):
                manifest = REVTOOL.extract(source, output, pdf_csv)
                repeated = REVTOOL.extract(source, output, pdf_csv)
            self.assertEqual(manifest, repeated)
            self.assertEqual(source.read_bytes(), before)
            self.assertEqual(manifest["sha256"], sha256(before).hexdigest())
            hidden = manifest["sheets"][1]
            self.assertEqual(hidden["state"], "hidden")
            self.assertEqual(hidden["formulas"], 1)
            self.assertEqual(hidden["missing_formula_caches"], 1)
            cells = json.loads((output / hidden["cells"]).read_text())
            formula = next(cell for cell in cells if cell["cell"] == "C2")
            self.assertEqual(formula["formula"], "=B2*2")
            self.assertIsNone(formula["value"])
            self.assertEqual(manifest["comparison_status_counts"], {"match": 1})
            for artifact in manifest["artifacts"]:
                self.assertEqual(sha256((output / artifact["file"]).read_bytes()).hexdigest(),
                                 artifact["sha256"])

    def test_package_paths_cannot_escape_extraction_directory(self):
        with tempfile.TemporaryDirectory() as tmp:
            root = Path(tmp)
            source = root / "fixture.xlsm"
            with ZipFile(source, "w") as archive:
                archive.writestr("customXml/../../../escaped.xml", "<test/>")
            with patch.object(REVTOOL, "VBA_Parser") as parser:
                parser.return_value.extract_macros.return_value = []
                with self.assertRaisesRegex(ValueError, "Unsafe workbook package path"):
                    REVTOOL.extract_code(source, root / "output")
            self.assertFalse((root / "escaped.xml").exists())


if __name__ == "__main__":
    unittest.main()
